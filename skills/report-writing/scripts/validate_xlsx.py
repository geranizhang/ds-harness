#!/usr/bin/env python3
"""Delivery table check (xlsx / csv, deterministic). report-writer runs it on every table it produces.

Usage: validate_xlsx.py <file.xlsx|file.csv> [...]

xlsx:
  1 no string cell starts with "=" (openpyxl stores it as a formula; Excel shows a repair dialog)
  2 columns whose header looks like a rate (rate/ratio/share/% or the adapter's words): numeric cells with a % number_format,
    never text, never a bare 0.186
  3 if a region column exists (names from adapter.json region_column_names): it contains the aggregate row label
    and one row per region in adapter.json regions (a missing region must be explained in the table note)
  4 last row of every sheet carries a provenance note (definition summary + per-region query/file source)
csv:
  5 if it contains non-ASCII text, it starts with a UTF-8 BOM (Excel otherwise mis-decodes it)
Exit: 0 / 1 / 3 missing openpyxl
"""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "_shared"))
import harness_config as hc  # noqa: E402

RATE_HDR = re.compile(r"率|占比|比例|%|pct|ratio|rate|share", re.I)
PROVENANCE = re.compile(r"口径|备注|query|来源|source|definition|note|provenance", re.I)


def check_xlsx(path, cfg):
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl missing (pip3 install openpyxl)")
        sys.exit(3)
    region_cols = {c.lower() for c in hc.get(cfg, "delivery.tables.region_column_names", ["region"])}
    agg = hc.get(cfg, "delivery.tables.aggregate_row_label", "ALL")
    regions = set(hc.get(cfg, "platform.regions", []))
    errs = []
    for ws in openpyxl.load_workbook(path).worksheets:
        rows = list(ws.iter_rows())
        if not rows:
            continue
        hdr = [(c.value if c.value is not None else "") for c in rows[0]]
        for row in rows:
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    errs.append(f"[{ws.title}] {c.coordinate} text starts with '=': {c.value[:30]!r}")
        for j, h in enumerate(hdr):
            if isinstance(h, str) and RATE_HDR.search(h):
                for row in rows[1:-1]:
                    c = row[j] if j < len(row) else None
                    if c is None or c.value in (None, ""):
                        continue
                    if isinstance(c.value, str):
                        errs.append(f"[{ws.title}] {c.coordinate} rate column '{h}' stored as text: {c.value!r}")
                    elif "%" not in str(c.number_format):
                        errs.append(f"[{ws.title}] {c.coordinate} rate column '{h}' has no % number_format (is {c.number_format})")
        for j, h in enumerate(hdr):
            if str(h).strip().lower() in region_cols:
                vals = {str(row[j].value).strip() for row in rows[1:] if j < len(row) and row[j].value is not None}
                if agg not in vals:
                    errs.append(f"[{ws.title}] region column lacks the '{agg}' aggregate row (has {sorted(vals)})")
                missing = regions - vals
                if missing:
                    errs.append(f"[{ws.title}] region column lacks rows for {sorted(missing)} (default is every region + {agg}; explain a missing region in the table note)")
        last = " ".join(str(c.value) for c in rows[-1] if c.value is not None)
        if not PROVENANCE.search(last):
            errs.append(f"[{ws.title}] last row has no provenance note (definition summary + per-region query/file source)")
    return errs


def check_csv(path):
    b = open(path, "rb").read()
    try:
        txt = b.decode("utf-8")
    except UnicodeDecodeError:
        return ["not valid UTF-8"]
    if any(ord(ch) > 127 for ch in txt) and not b.startswith(b"\xef\xbb\xbf"):
        return ["contains non-ASCII text but no UTF-8 BOM; Excel will mis-decode it (write with encoding='utf-8-sig')"]
    return []


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    cfg = hc.load()
    rc = 0
    for p in sys.argv[1:]:
        errs = check_xlsx(p, cfg) if p.lower().endswith((".xlsx", ".xlsm")) else check_csv(p)
        print(f"== {p} {hc.describe(cfg)} ==")
        for e in errs:
            print("ERROR:", e)
        if not errs:
            print("OK")
        else:
            rc = 1
    sys.exit(rc)


if __name__ == "__main__":
    main()
